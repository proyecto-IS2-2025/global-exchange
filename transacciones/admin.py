from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone
from .models import Transaccion, HistorialTransaccion, ConfiguracionTransaccion
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from decimal import Decimal


def exportar_transacciones_excel(modeladmin, request, queryset):
    """
    Acción del admin para exportar transacciones seleccionadas a Excel
    """
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Transacciones"
    
    # Estilos
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título
    ws.merge_cells('A1:M1')
    title_cell = ws['A1']
    title_cell.value = 'REPORTE DE TRANSACCIONES'
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Subtítulo con fecha
    ws.merge_cells('A2:M2')
    subtitle_cell = ws['A2']
    subtitle_cell.value = f'Generado el {timezone.now().strftime("%d/%m/%Y %H:%M")}'
    subtitle_cell.alignment = Alignment(horizontal='center')
    
    # Encabezados
    headers = [
        'Nº Transacción',
        'Fecha',
        'Cliente',
        'Tipo Operación',
        'Estado',
        'Divisa Origen',
        'Monto Origen',
        'Divisa Destino',
        'Monto Destino',
        'Tasa Aplicada',
        'Código TAUSER',
        'Terminal',
        'Procesado Por'
    ]
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Datos
    row = 4
    for transaccion in queryset.select_related('cliente', 'divisa_origen', 'divisa_destino', 'tauser_terminal', 'tauser_deposito', 'procesado_por'):
        ws.cell(row=row, column=1, value=transaccion.numero_transaccion)
        ws.cell(row=row, column=2, value=transaccion.fecha_creacion.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row, column=3, value=transaccion.cliente.nombre_completo if transaccion.cliente else 'N/A')
        ws.cell(row=row, column=4, value=transaccion.tipo_operacion.upper())
        ws.cell(row=row, column=5, value=transaccion.estado.upper())
        ws.cell(row=row, column=6, value=transaccion.divisa_origen.code if transaccion.divisa_origen else 'N/A')
        ws.cell(row=row, column=7, value=float(transaccion.monto_origen) if transaccion.monto_origen else 0)
        ws.cell(row=row, column=8, value=transaccion.divisa_destino.code if transaccion.divisa_destino else 'N/A')
        ws.cell(row=row, column=9, value=float(transaccion.monto_destino) if transaccion.monto_destino else 0)
        ws.cell(row=row, column=10, value=float(transaccion.tasa_de_cambio_aplicada) if transaccion.tasa_de_cambio_aplicada else 0)
        ws.cell(row=row, column=11, value=transaccion.tauser_code or 'N/A')
        
        # Terminal: compras usan tauser_terminal, ventas usan tauser_deposito
        if transaccion.tipo_operacion == 'compra':
            terminal_value = str(transaccion.tauser_terminal) if transaccion.tauser_terminal else 'N/A'
        else:  # venta
            terminal_value = str(transaccion.tauser_deposito) if transaccion.tauser_deposito else 'N/A'
        ws.cell(row=row, column=12, value=terminal_value)
        
        ws.cell(row=row, column=13, value=transaccion.procesado_por.get_full_name() if transaccion.procesado_por else 'N/A')
        
        # Aplicar bordes
        for col in range(1, 14):
            ws.cell(row=row, column=col).border = border
        
        # Formato de números para columnas de montos
        ws.cell(row=row, column=7).number_format = '#,##0.00'
        ws.cell(row=row, column=9).number_format = '#,##0.00'
        ws.cell(row=row, column=10).number_format = '#,##0.00'
        
        row += 1
    
    # Ajustar ancho de columnas
    column_widths = {
        'A': 18, 'B': 16, 'C': 25, 'D': 15, 'E': 12,
        'F': 12, 'G': 15, 'H': 12, 'I': 15, 'J': 12,
        'K': 15, 'L': 20, 'M': 20
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Crear respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'transacciones_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename={filename}'
    
    wb.save(response)
    return response

exportar_transacciones_excel.short_description = "Exportar transacciones seleccionadas a Excel"


@admin.register(Transaccion)
class TransaccionAdmin(admin.ModelAdmin):
    list_display = ('numero_transaccion', 'tipo_operacion', 'cliente', 'estado', 'tauser_code', 'tauser_terminal', 'monto_origen', 'monto_destino', 'fecha_creacion')
    list_filter = ('tipo_operacion', 'estado', 'fecha_creacion')
    search_fields = ('numero_transaccion', 'tauser_code', 'cliente__nombre', 'cliente__apellido')
    readonly_fields = ('numero_transaccion', 'tauser_code', 'fecha_creacion', 'fecha_actualizacion')
    list_select_related = ('cliente', 'tauser_terminal', 'divisa_origen', 'divisa_destino')
    actions = [exportar_transacciones_excel]
    
    def get_fieldsets(self, request, obj=None):
        """
        Retorna fieldsets dinámicos según si el usuario es staff/admin
        """
        fieldsets = [
            ('Información Básica', {
                'fields': ('numero_transaccion', 'tipo_operacion', 'estado', 'cliente')
            }),
            ('Divisas y Montos', {
                'fields': ('divisa_origen', 'monto_origen', 'divisa_destino', 'monto_destino', 'tasa_de_cambio_aplicada')
            }),
            ('TAUSER', {
                'fields': ('tauser_code', 'tauser_terminal', 'tauser_deposito'),
                'description': 'TAUSER de retiro (compras) y TAUSER de depósito (ventas)'
            }),
            ('Medio de Pago', {
                'fields': ('medio_pago_datos',)
            }),
            ('Información Adicional', {
                'fields': ('observaciones', 'observacion', 'procesado_por', 'fecha_creacion', 'fecha_actualizacion')
            }),
        ]
        
        # Solo mostrar análisis de ganancias a administradores y superusuarios
        if request.user.is_superuser or request.user.groups.filter(name='admin').exists():
            fieldsets.insert(3, (
                'Análisis de Ganancias (Solo Administradores)', {
                    'fields': ('tasa_base', 'margen_spread', 'porcentaje_comision', 'comision_aplicada'),
                    'classes': ('collapse',),
                    'description': 'Información detallada de costos y márgenes para análisis de rentabilidad'
                }
            ))
        
        return fieldsets


@admin.register(HistorialTransaccion)
class HistorialTransaccionAdmin(admin.ModelAdmin):
    list_display = ('transaccion', 'estado_anterior', 'estado_nuevo', 'fecha_cambio', 'modificado_por')
    list_filter = ('estado_nuevo', 'fecha_cambio')
    search_fields = ('transaccion__numero_transaccion',)
    readonly_fields = ('fecha_cambio',)


@admin.register(ConfiguracionTransaccion)
class ConfiguracionTransaccionAdmin(admin.ModelAdmin):
    list_display = ('nombre', 'valor', 'fecha_modificacion', 'modificado_por')
    search_fields = ('nombre', 'descripcion')
    readonly_fields = ('fecha_modificacion',)
